"""Чтение списка проектов из XLSX.

`data_only=True` обязателен: без него ячейка с формулой отдаёт саму формулу
(`=CONCAT(...)`), и домен превращается в строку, которую нормализация честно
забракует — а причина будет выглядеть как ошибка отдела, а не как наша.

Значения приводятся к строкам здесь, в одном месте: Excel хранит даты числами, а
объём работ — float'ом, и `120.0` вместо `120` дошло бы до отчёта.
"""

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
from typing import IO, Any

from openpyxl import load_workbook

from ahrefs_cases.intake.rows import RawTable, table_from_matrix


def read_xlsx(path: Path) -> RawTable:
    """Книга на диске → сырая таблица. Имя файла становится происхождением."""
    with path.open("rb") as stream:
        return read_xlsx_stream(stream, origin=str(path))


def read_xlsx_stream(stream: IO[bytes], origin: str) -> RawTable:
    """Первый лист книги → сырая таблица.

    Читаем именно первый лист, а не лист по имени: имя у каждого отдела своё
    («Лист1», «домены», «Sheet1»), а порядок один.

    Поток, а не только путь: книга приходит и телом HTTP-запроса (экран
    загрузки Ф6), и писать её на диск ради `openpyxl` значило бы заводить
    временные файлы и убирать их за собой — при том, что читатель и так
    работает с потоком.
    """
    workbook = load_workbook(filename=stream, read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        rows = [[_cell_to_str(cell) for cell in row] for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()

    return table_from_matrix(origin, rows)


def _cell_to_str(value: Any) -> str:
    """Ячейка → строка в том виде, в каком её ждёт валидация.

    Дата приводится к ISO, целое число без дробной части — к целому. Иначе
    `2025-01-01 00:00:00` и `120.0` уехали бы в правила, где их пришлось бы
    разбирать второй раз, уже не зная, что это пришло из Excel.
    """
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()
